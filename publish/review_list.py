#!/usr/bin/env python3
"""Reviewer worklist of every [?] in a volume.  python3 review_list.py 2   (needs work/volN_clean.md from the cleaners)"""
import re, sys, pathlib, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

PUB = pathlib.Path(__file__).resolve().parent
DEV = re.compile(r'[ऀ-ॿ]')
IAST = re.compile(r'[āīūṛṝḷṃḥśṣṭḍṇñṅ]', re.I)
REF = re.compile(r'([0-9\u0966-\u096F]|P\.|Pā\.|\u092a\u093e\.|\u0909\.|\u090b\.|Uṇ|Nir|Ṛ\.|Ṛg|Śat|Ait|Tai|tai\.|sūtra|Phiṭ|Aṣṭ|vārttika|vārtika|Paribhāṣā|ā\. ?\[|varga|Varga)', re.I)
LAYERS = [('Saṃhitā-pāṭh', 'Saṃhitā text'), ('Pada-pāṭh', 'Pada text'), ('Sāyaṇa-bhāṣya', 'Sāyaṇa-bhāṣya'),
          ('Prati-padārth', 'Pratipadārtha'), ('Bhāvārth', 'Bhāvārtha'), ('English', 'English'),
          ('Viśeṣa', 'Special Topics'), ('Vyākaraṇa', 'Grammar'), ('Grammar', 'Grammar'), ('Anuvāda', 'Kannada rendering')]

def clean_ctx(s):
    s = re.sub(r'<[^>]+>', '', s)
    s = re.sub(r'[*`>]+', '', s)
    return re.sub(r'\s+', ' ', s).strip()

def collect(md):
    rows, sukta, head, page, layer = [], '(front matter)', '', '', 'Text'
    for line in md.split('\n'):
        m = re.match(r'^## .*S[ūu]kta\s+(\d+)', line)
        if m: sukta = f'Sūkta {m.group(1)}'; head = ''; layer = 'Text'
        elif line.startswith('## '): sukta = re.sub(r'[*]', '', line[3:])[:60]; head = ''
        elif line.startswith('# '): sukta = re.sub(r'[*]', '', line[2:])[:60]
        elif line.startswith('### '): head = re.sub(r'[*]', '', line[4:]).strip()[:90]
        m = re.search(r'original pp?\. ([\d–ivxlc]+)', line) or re.match(r'^#{2,3} \[?Page ([\d–ivxlc]+)', line)
        if m: page = m.group(1)
        if line.startswith('**॥') or line.startswith('**Grammar') or line.startswith('**Translation') or line.startswith('**Pratipad'):
            for key, lab in LAYERS:
                if key.lower() in line[:80].lower(): layer = lab; break
        if '[?]' not in line: continue
        in_note = line.lstrip().startswith('*(')
        quote = line.startswith('>')
        for m in re.finditer(r'\[\?\]', line):
            before = clean_ctx(line[max(0, m.start() - 160):m.start()])
            after = clean_ctx(line[m.end():m.end() + 100])
            near = before[-34:]
            if REF.search(near[-16:]): cat = 'Reference / numeral'
            elif DEV.search(near) or IAST.search(near[-22:]): cat = 'Sanskrit word / phrase'
            else: cat = 'Sense / name / other'
            lyr = 'Note' if in_note else layer
            if cat == 'Sanskrit word / phrase' and (quote or lyr in ('Saṃhitā text', 'Pada text', 'Sāyaṇa-bhāṣya')): pri = 'High'
            elif cat == 'Reference / numeral': pri = 'Low'
            else: pri = 'Medium'
            rows.append([sukta, head, page, lyr, cat, pri, before[-120:], '⟦?⟧', after, '', '', ''])
    return rows

def main(vol):
    md = (PUB / 'work' / f'vol{vol}_clean.md').read_text(encoding='utf-8')
    rows = collect(md)
    for i, r in enumerate(rows, 1): r.insert(0, i)
    wb = Workbook(); ws = wb.active; ws.title = 'Worklist'
    hdr = ['ID', 'Sūkta / section', 'Rik / heading', 'Original page', 'Layer', 'Category', 'Priority',
           'Text before the flag', 'Flag', 'Text after the flag', "Reviewer's reading", 'Decision (confirm / correct / illegible)', 'Comment']
    ws.append(hdr)
    for r in rows: ws.append(r)
    fill = PatternFill('solid', fgColor='DDDDDD')
    for c in ws[1]: c.font = Font(bold=True); c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='top')
    widths = [6, 18, 30, 9, 16, 20, 9, 55, 6, 40, 28, 22, 28]
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
        row[8].font = Font(bold=True, color='C00000')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    # summary
    s = wb.create_sheet('Summary')
    tot = collections.Counter(r[6] for r in rows)
    s.append([f'Volume {vol}: reviewer worklist — {len(rows)} flagged spots']); s['A1'].font = Font(bold=True, size=13)
    s.append([]); s.append(['By priority', 'Count'])
    for k in ('High', 'Medium', 'Low'): s.append([k, tot.get(k, 0)])
    s.append([]); s.append(['By category', 'Count'])
    for k, v in collections.Counter(r[5] for r in rows).most_common(): s.append([k, v])
    s.append([]); s.append(['By layer', 'Count'])
    for k, v in collections.Counter(r[4] for r in rows).most_common(): s.append([k, v])
    s.append([]); s.append(['By Sūkta / section', 'Total', 'High'])
    by = collections.OrderedDict()
    for r in rows: by.setdefault(r[1], [0, 0]); by[r[1]][0] += 1; by[r[1]][1] += (r[6] == 'High')
    for k, (a, b) in by.items(): s.append([k, a, b])
    s.column_dimensions['A'].width = 38
    # how to use
    h = wb.create_sheet('How to use')
    for t in ["Each row is one place in the translation marked [?] = a reading that could not be made out with confidence from the 1949 scan.",
              "Original page = the page number of the 1949 Kannada edition (the translation keeps these as 'original p. N' markers).",
              "Priority High = a doubtful Sanskrit word or phrase in the Saṃhitā/Pada/bhāṣya text (affects the text itself).",
              "Priority Medium = a doubtful word, name or sense elsewhere. Priority Low = a doubtful reference or sūtra number.",
              "For each row, check the original page, then fill 'Reviewer's reading' (what the print actually shows) and 'Decision':",
              "   confirm = the translation's reading is right; correct = give the right reading; illegible = cannot be made out even in the original.",
              "Corrections returned in this sheet can be applied to the text and the PDF rebuilt."]:
        h.append([t])
    h.column_dimensions['A'].width = 130
    out = PUB / 'review' / f'Vol{vol}_review_worklist.xlsx'; wb.save(out)
    print(f'wrote {out}: {len(rows)} spots; priority {dict(tot)}')
    print('by category:', dict(collections.Counter(r[5] for r in rows)))
    return rows

if __name__ == '__main__': main(int(sys.argv[1]))
