#!/usr/bin/env python3
"""One reviewer worklist for all six volumes (every [?]).  python3 review_all.py"""
import re, pathlib, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import review_list as R
PUB = R.PUB; ROOT = PUB.parent
SRC = {1: PUB/'work'/'vol1_clean.md', 2: PUB/'work'/'vol2_clean.md', 3: ROOT/'Rigveda_Samhita_Vol3_English_Translation.md',
       4: PUB/'work'/'vol4_clean.md', 5: PUB/'work'/'vol5_clean.md',
       6: ROOT/'Rigveda_Samhita_Vol6_English_Translation.md'}
allrows = []
for v, p in SRC.items():
    rows = R.collect(p.read_text(encoding='utf-8'))
    for r in rows: allrows.append([f'Vol {v}'] + r)
for i, r in enumerate(allrows, 1): r.insert(0, i)
wb = Workbook(); ws = wb.active; ws.title = 'Worklist'
hdr = ['ID','Volume','Sūkta / section','Rik / heading','Original page','Layer','Category','Priority','Text before the flag','Flag','Text after the flag',"Reviewer's reading",'Decision (confirm / correct / illegible)','Comment']
ws.append(hdr)
for r in allrows: ws.append(r)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDDDDD'); c.alignment = Alignment(wrap_text=True, vertical='top')
for i, w in enumerate([7,7,18,30,9,16,20,9,55,6,40,28,22,28], 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    row[9].font = Font(bold=True, color='C00000')
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
s = wb.create_sheet('Summary')
s.append([f'All volumes: reviewer worklist — {len(allrows)} flagged spots']); s['A1'].font = Font(bold=True, size=13)
s.append([]); s.append(['Volume','Total','High','Medium','Low'])
for v in SRC:
    rr = [r for r in allrows if r[1] == f'Vol {v}']; c = collections.Counter(r[7] for r in rr)
    s.append([f'Vol {v}', len(rr), c['High'], c['Medium'], c['Low']])
c = collections.Counter(r[7] for r in allrows); s.append(['All', len(allrows), c['High'], c['Medium'], c['Low']])
s.append([]); s.append(['By category','Count'])
for k, n in collections.Counter(r[6] for r in allrows).most_common(): s.append([k, n])
s.append([]); s.append(['By layer','Count'])
for k, n in collections.Counter(r[5] for r in allrows).most_common(): s.append([k, n])
s.column_dimensions['A'].width = 38
h = wb.create_sheet('How to use')
for t in ["Each row is one place in the translation marked [?] = a reading that could not be made out with confidence from the 1949 scan.",
          "Volume 1 includes the portions translated late (printed pp. 95–167, 191–205, 261–283, 288–295). Volumes 3 and 6 are read from their closed source files (Volume 6 is the uncleaned working text: its page headings and progress note are not stripped).",
          "Original page = the page number of the 1949 Kannada edition.",
          "Priority High = a doubtful Sanskrit word or phrase in the Saṃhitā/Pada/bhāṣya text. Medium = a doubtful word, name or sense elsewhere. Low = a doubtful reference or sūtra number.",
          "Not flagged with [?] but also needing review: Vol 1 rishi-index hymn references (Kannada numerals, not reproduced); Vol 1 tables on pp. 117–118, 138–139, 150–151 (not reproduced); Vol 1 Pada-pāṭha tables pp. 291–294 (accents not reproduced).",
          "Fill 'Reviewer's reading' and 'Decision' (confirm / correct / illegible); corrections can be applied and the PDFs rebuilt."]:
    h.append([t])
h.column_dimensions['A'].width = 140
extra = wb.create_sheet('Other review points')
extra.append(['Volume','Printed pages','Point','Why it needs review'])
for row in [
 ['Vol 1','191–205','Ṛṣi index: hymn references (Kannada numerals) omitted','Illegible at 150 dpi; names and verse-counts only. Needs a zoomed re-reading.'],
 ['Vol 1','117–118','Per-adhyāya Mādhyandina table not reproduced','Digits unreadable; only verified rows and totals given.'],
 ['Vol 1','138–139','Taittirīya Saṃhitā/Brāhmaṇa prapāṭhaka tables not reproduced','Only praśna counts and printed totals given.'],
 ['Vol 1','150–151','Jaiminīya Sāma counts not reproduced','Digits unreadable.'],
 ['Vol 1','140–143','Branch numbering 30–41 unreliable','Headings read at low resolution.'],
 ['Vol 1','145','"13 ācāryas" count vs 14 names as read','Source count may be misread.'],
 ['Vol 1','264','Colophon weekday: Sanskrit "bhaume" (Tuesday) vs Kannada "budhavāra" (Wednesday)','Inconsistency in the source, left as printed.'],
 ['Vol 1','291–294','Pada-pāṭha accent marks not reproduced','Notation not converted.'],
 ['Vol 1','261–283','Dates/folio numerals for several commentators (Kālanātha, Kṣura, Śobhākara, Sāyaṇa dates)','Kannada digits uncertain.'],
 ['Vol 6','521, 512–518, 525, 545–546, 551, 579, 587, 590, 594, 611, 617','Garbled Ṛgvedic citations and the compressed kārīryām passage (p. 521) given as read','Verse text unclear at 150 dpi; glosses are tentative. Needs zoomed re-reading against the Ṛgveda text.'],
 ['Vol 6','525, 530, 538, 569','Yāska/Nirukta lines and the varṣitā quotation given as read','Reading and Nirukta references uncertain.'],
 ['Vol 6','471–472, 582, 589, 595','Grammar page of Rik 76.1 folded across two pages; accent arguments in 80.6, 80.8, 80.12, 80.13 and a compressed clause in 80.9 read with doubt','Compressed scholastic argument; sūtra numerals tentative throughout.'],
 ['Vol 6','615–618','Dadhyañc/Manu reference lists and the 9×3×3×10=810 computation (Kannada numerals)','Digits read as given; not verified.'],
 ['Vol 6','all','Varga numerals in headings and source cross-reference "Ṛk Saṃhitā vol. N pp. 640–641" left unreconciled','Small Kannada numerals unreliable.'],
 ['All','—','Outside Vedic/Sanskrit expert review of all [?] readings','Recommended before final publication.'],
 ['All','—','Present Maharaja/Guruji portraits, captions, copyright/permission details','Still to be supplied.']]:
    extra.append(row)
for i, w in enumerate([9,14,70,70], 1): extra.column_dimensions[get_column_letter(i)].width = w
for row in extra.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
for c in extra[1]: c.font = Font(bold=True)
out = PUB/'review'/'All_volumes_review_worklist.xlsx'; wb.save(out)
print(out, len(allrows)); print([(v, sum(1 for r in allrows if r[1]==f'Vol {v}')) for v in SRC])
